# create_testing_doc.py
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create document
doc = Document()

# Add title
title = doc.add_heading("Loan Management System - Users API Testing Guide", 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add subtitle
subtitle = doc.add_paragraph("Comprehensive Testing Documentation")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].bold = True

# Add metadata table
table = doc.add_table(rows=4, cols=2)
table.style = "LightShading-Accent1"
table.cell(0, 0).text = "Document Version:"
table.cell(0, 1).text = "1.0"
table.cell(1, 0).text = "Date:"
table.cell(1, 1).text = "December 2025"
table.cell(2, 0).text = "Author:"
table.cell(2, 1).text = "Loan Management System Team"
table.cell(3, 0).text = "Project:"
table.cell(3, 1).text = "LMS Backend API"

doc.add_paragraph()  # Add space

# Section 1: Introduction
doc.add_heading("1. Introduction", level=1)
intro = doc.add_paragraph()
intro.add_run(
    "This document provides comprehensive testing instructions for the Users API of the Loan Management System. "
)
intro.add_run(
    "The API includes authentication endpoints for user registration, login, profile management, and token handling."
)

# Section 2: Testing Environment
doc.add_heading("2. Testing Environment Setup", level=1)

doc.add_heading("2.1 Prerequisites", level=2)
prereq = doc.add_paragraph("Ensure the following are available:")
doc.add_paragraph("✅ Django development server running", style="List Bullet")
doc.add_paragraph("✅ PostgreSQL database configured", style="List Bullet")
doc.add_paragraph("✅ Browser with internet access", style="List Bullet")
doc.add_paragraph("✅ API documentation accessible", style="List Bullet")

doc.add_heading("2.2 URLs", level=2)
url_table = doc.add_table(rows=3, cols=2)
url_table.style = "LightGrid-Accent1"
url_table.cell(0, 0).text = "Service"
url_table.cell(0, 1).text = "URL"
url_table.cell(1, 0).text = "API Documentation"
url_table.cell(1, 1).text = "http://localhost:8000/api/docs/"
url_table.cell(2, 0).text = "Django Admin"
url_table.cell(2, 1).text = "http://localhost:8000/admin/"

# Section 3: API Endpoints
doc.add_heading("3. API Endpoints Overview", level=1)
endpoints_table = doc.add_table(rows=7, cols=4)
endpoints_table.style = "MediumGrid1-Accent1"
headers = ["Endpoint", "Method", "Description", "Auth Required"]
for i, header in enumerate(headers):
    endpoints_table.cell(0, i).text = header
    endpoints_table.cell(0, i).paragraphs[0].runs[0].bold = True

data = [
    ["/register/", "POST", "Register new user", "No"],
    ["/login/", "POST", "User login", "No"],
    ["/profile/", "GET", "Get user profile", "Yes"],
    ["/profile/", "PUT/PATCH", "Update user profile", "Yes"],
    ["/token/refresh/", "POST", "Refresh access token", "No"],
    ["/logout/", "POST", "User logout", "Yes"],
]

for row_idx, row_data in enumerate(data, 1):
    for col_idx, cell_data in enumerate(row_data):
        endpoints_table.cell(row_idx, col_idx).text = cell_data

# Save document
filename = "LMS_API_Testing_Guide.docx"
doc.save(filename)
print(f"✅ Word document created: {filename}")
